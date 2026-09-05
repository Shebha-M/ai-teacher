"""PowerPoint PPTX extractor leveraging python-pptx for slide-by-slide text extraction."""
import io
import logging
from typing import List
import pptx
from ai_engine.rag.extractors.base import BaseExtractor, ExtractedDocument, ExtractedSection

logger = logging.getLogger("ai_teacher.rag.pptx")


class PPTXExtractor(BaseExtractor):
    """Extracts slide text, titles, numbers, and notes from PowerPoint presentations."""

    def extract(self, file_bytes: bytes, filename: str) -> ExtractedDocument:
        try:
            prs = pptx.Presentation(io.BytesIO(file_bytes))
            total_slides = len(prs.slides)
            sections: List[ExtractedSection] = []
            full_text_parts: List[str] = []

            for idx, slide in enumerate(prs.slides):
                slide_num = idx + 1
                slide_title = f"Slide {slide_num}"
                text_items: List[str] = []

                # Find slide title shape if available
                if slide.shapes.title and slide.shapes.title.text.strip():
                    slide_title = slide.shapes.title.text.strip()

                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for paragraph in shape.text_frame.paragraphs:
                            line = paragraph.text.strip()
                            if line and line != slide_title:
                                text_items.append(line)

                # Check for speaker notes
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                    notes = slide.notes_slide.notes_text_frame.text.strip()
                    if notes:
                        text_items.append(f"[Speaker Notes: {notes}]")

                content_str = "\n".join(text_items)
                if content_str or slide_title:
                    sections.append(
                        ExtractedSection(
                            title=slide_title,
                            content=content_str,
                            slide_number=slide_num,
                            metadata={"slide": slide_num, "title": slide_title},
                        )
                    )
                    full_text_parts.append(f"### Slide {slide_num}: {slide_title}\n{content_str}")

            doc_metadata = {}
            if prs.core_properties:
                doc_metadata = {
                    "title": prs.core_properties.title or None,
                    "author": prs.core_properties.author or None,
                }

            return ExtractedDocument(
                filename=filename,
                file_type="pptx",
                total_pages_or_slides=total_slides,
                full_text="\n\n".join(full_text_parts),
                sections=sections,
                metadata=doc_metadata,
            )
        except Exception as e:
            logger.error(f"Error extracting PPTX {filename}: {e}", exc_info=True)
            raise ValueError(f"Failed to process PowerPoint presentation '{filename}': {str(e)}") from e
